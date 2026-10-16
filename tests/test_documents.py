"""PDF and DOCX ingest (the `pdf` and `docx` extras).

PDFs are built byte by byte by `make_pdf` (text on each page via the
standard Helvetica font), DOCX files with python-docx, so no fixture is
downloaded or committed. Tests that need an extra skip without it; the
"not installed" tests simulate a failing import instead.
"""

from __future__ import annotations

import contextlib
import io
import sys
from pathlib import Path
from typing import Any

import pytest
from mcp import Client
from mcp_types import TextContent
from typer.testing import CliRunner

from notelight import documents
from notelight.cli import app
from notelight.embed import HashingEmbedder
from notelight.indexer import sync
from notelight.ingest import walk
from notelight.mcp_server import NotesService, create_server
from notelight.search import Searcher, check_staleness, skipped_files

PDF_PAGES = [
    ["Quarterly report", "Revenue grew in every region."],
    ["Glaciers retreat as summers warm.", "Meltwater feeds the rivers."],
    [],  # a blank page between two pages of text
    ["Appendix: sourdough needs a starter."],
]


# -- fixture builders ------------------------------------------------------------
def _pdf_string(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal valid PDF with one text line per entry on each page."""
    page_ids = [4 + 2 * i for i in range(len(pages))]
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_id, lines in zip(page_ids, pages, strict=True):
        shows = " ".join(f"({_pdf_string(line)}) Tj T*" for line in lines)
        content = f"BT /F1 12 Tf 16 TL 72 720 Td {shows} ET".encode() if lines else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        objects.append(
            f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream"
        )
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(out)


def encrypt_pdf(data: bytes, user_password: str) -> bytes:
    pypdf = pytest.importorskip("pypdf")
    writer = pypdf.PdfWriter(clone_from=pypdf.PdfReader(io.BytesIO(data)))
    writer.encrypt(user_password=user_password, owner_password="owner", algorithm="RC4-128")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_docx(path: Path) -> None:
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("Field notes from the trip.")
    document.add_heading("Mountains", level=1)
    document.add_paragraph("Glaciers retreat as summers warm.")
    document.add_heading("Huts", level=2)
    document.add_paragraph("# not a heading, just a hash")
    table = document.add_table(rows=2, cols=2)
    for r, row in enumerate([["Hut", "Altitude"], ["Refuge", "2800 m"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_heading("Food", level=1)
    document.add_paragraph("Sourdough keeps for a week.")
    document.save(str(path))


@pytest.fixture
def pdf_extra() -> None:
    pytest.importorskip("pypdf")


@pytest.fixture
def docx_extra() -> None:
    pytest.importorskip("docx")


@pytest.fixture
def library(tmp_path: Path) -> Path:
    """A folder with one note, one PDF and one DOCX (if python-docx is installed)."""
    root = tmp_path / "library"
    root.mkdir()
    (root / "note.md").write_text("# Note\n\nKites fly in wind.\n", encoding="utf-8")
    (root / "report.pdf").write_bytes(make_pdf(PDF_PAGES))
    with contextlib.suppress(pytest.skip.Exception):
        make_docx(root / "trip.docx")
    return root


@pytest.fixture
def without_extras(monkeypatch: pytest.MonkeyPatch) -> None:
    """Make `import pypdf` and `import docx` fail, as if the extras were absent."""
    for name in ["pypdf", "docx"]:
        monkeypatch.setitem(sys.modules, name, None)


# -- extraction ------------------------------------------------------------------
def test_pdf_pages_are_separated_by_form_feeds(pdf_extra: None) -> None:
    decoded = documents.extract(".pdf", make_pdf(PDF_PAGES))
    assert decoded.text is not None and decoded.skip_reason is None
    pages = decoded.text.split(documents.PAGE_BREAK)
    assert len(pages) == 4
    assert "Revenue grew" in pages[0] and "Glaciers" in pages[1]
    assert pages[2] == "\n"  # the blank page keeps its own line
    assert all(p.endswith("\n") for p in pages)
    lines = decoded.text.split("\n")
    starts = documents.page_starts(decoded.text)
    assert len(starts) == 4 and starts[0] == 1
    # Each page after the first starts on the line holding its form feed.
    assert all(lines[s - 1].startswith(documents.PAGE_BREAK) for s in starts[1:])
    glacier = next(i for i, line in enumerate(lines, 1) if "Glaciers" in line)
    assert documents.page_range(starts, glacier, glacier) == (2, 2)


def test_page_range() -> None:
    starts = [1, 4, 9]
    assert documents.page_range(starts, 1, 3) == (1, 1)
    assert documents.page_range(starts, 3, 4) == (1, 2)
    assert documents.page_range(starts, 5, 20) == (2, 3)
    assert documents.page_starts("one page\nonly\n") == [1]
    assert documents.page_starts("a\n\fb\n\fc\n") == [1, 2, 3]


def test_pdf_with_user_password_is_skipped(pdf_extra: None) -> None:
    decoded = documents.extract(".pdf", encrypt_pdf(make_pdf(PDF_PAGES), "secret"))
    assert decoded.text is None
    assert decoded.skip_reason == "encrypted PDF (password required)"


def test_pdf_with_only_an_owner_password_is_read(pdf_extra: None) -> None:
    decoded = documents.extract(".pdf", encrypt_pdf(make_pdf(PDF_PAGES), ""))
    assert decoded.text is not None and "Glaciers" in decoded.text


def test_image_only_pdf_is_skipped(pdf_extra: None) -> None:
    decoded = documents.extract(".pdf", make_pdf([[], []]))
    assert decoded.text is None
    assert decoded.skip_reason is not None and "no text layer" in decoded.skip_reason


@pytest.mark.parametrize("data", [b"", b"%PDF-1.4\ngarbage", b"not a pdf at all"])
def test_malformed_pdf_is_skipped_not_raised(pdf_extra: None, data: bytes) -> None:
    decoded = documents.extract(".pdf", data)
    assert decoded.text is None
    assert decoded.skip_reason is not None
    assert decoded.skip_reason.startswith(("unreadable PDF", "PDF has no text"))


def test_docx_headings_paragraphs_and_tables(docx_extra: None, tmp_path: Path) -> None:
    make_docx(tmp_path / "trip.docx")
    decoded = documents.extract(".docx", (tmp_path / "trip.docx").read_bytes())
    assert decoded.text == (
        "Field notes from the trip.\n\n"
        "# Mountains\n\n"
        "Glaciers retreat as summers warm.\n\n"
        "## Huts\n\n"
        "\\# not a heading, just a hash\n\n"
        "Hut | Altitude\nRefuge | 2800 m\n\n"
        "# Food\n\n"
        "Sourdough keeps for a week.\n\n"
    )


def test_empty_and_malformed_docx_are_skipped(docx_extra: None, tmp_path: Path) -> None:
    import docx

    docx.Document().save(str(tmp_path / "empty.docx"))
    decoded = documents.extract(".docx", (tmp_path / "empty.docx").read_bytes())
    assert (decoded.text, decoded.skip_reason) == (None, "DOCX has no text")
    decoded = documents.extract(".docx", b"PK\x03\x04 not really a zip")
    assert decoded.skip_reason is not None and decoded.skip_reason.startswith("unreadable DOCX")


def test_extractor_id_names_library_and_version(pdf_extra: None, docx_extra: None) -> None:
    assert documents.extractor_id(".pdf").startswith(f"pdf{documents.EXTRACT_VERSION}:pypdf-")
    assert documents.extractor_id(".DOCX").startswith("docx1:python-docx-")
    assert documents.content_mode(".md") == "md"
    assert documents.content_mode(".txt") == "txt"


# -- the extras are optional ------------------------------------------------------
def test_missing_extra_reasons(without_extras: None) -> None:
    assert documents.missing_extra(".pdf") == (
        "PDF support is not installed; install notelight[pdf]"
    )
    assert documents.missing_extra(".docx") == (
        "DOCX support is not installed; install notelight[docx]"
    )
    assert documents.missing_extra(".md") is None
    assert documents.is_missing_extra_reason(documents.missing_extra(".pdf") or "")
    assert not documents.is_missing_extra_reason("binary")


def test_without_extras_documents_are_skipped_with_a_reason(
    library: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    (library / "trip.docx").write_bytes(b"PK\x03\x04")  # never read without the extra
    monkeypatch.setitem(sys.modules, "pypdf", None)
    monkeypatch.setitem(sys.modules, "docx", None)
    stats = sync(library, HashingEmbedder())
    assert stats.files == 1
    reasons = {s.path: s.reason for s in stats.skipped}
    assert reasons == {
        "report.pdf": "PDF support is not installed; install notelight[pdf]",
        "trip.docx": "DOCX support is not installed; install notelight[docx]",
    }
    assert set(skipped_files(library)) == {"report.pdf", "trip.docx"}
    assert not check_staleness(library).stale  # skipped is not "new"
    # A second sync reuses the recorded skip and still reports it.
    again = sync(library, HashingEmbedder())
    assert again.up_to_date
    assert {s.path for s in again.skipped} == {"report.pdf", "trip.docx"}


def test_installing_the_extra_later_indexes_skipped_files(
    library: Path, monkeypatch: pytest.MonkeyPatch, pdf_extra: None
) -> None:
    (library / "trip.docx").unlink(missing_ok=True)
    with monkeypatch.context() as m:
        m.setitem(sys.modules, "pypdf", None)
        assert sync(library, HashingEmbedder()).files == 1
    # The PDF itself did not change, but it can be read now.
    assert check_staleness(library).added == 1
    stats = sync(library, HashingEmbedder())
    assert (stats.files, stats.added, stats.skipped) == (2, 1, [])


# -- indexing, search and citations -----------------------------------------------
def test_walk_finds_documents(library: Path) -> None:
    (library / "slides.pptx").write_bytes(b"x")
    found = [f.rel_path for f in walk(library)]
    assert "report.pdf" in found and "slides.pptx" not in found


def test_pdf_hits_carry_pages(library: Path, pdf_extra: None) -> None:
    sync(library, HashingEmbedder())
    searcher = Searcher(library, HashingEmbedder())
    hit = next(r for r in searcher.search("glaciers retreat", k=5) if r.path == "report.pdf")
    assert hit.heading_path == []
    assert hit.location == f"report.pdf:{hit.line_start}-{hit.line_end}"
    assert hit.page_start is not None and hit.page_end is not None
    assert hit.page_start <= 2 <= hit.page_end
    assert hit.pages == (
        f"p. {hit.page_start}"
        if hit.page_start == hit.page_end
        else f"pp. {hit.page_start}-{hit.page_end}"
    )
    assert hit.to_dict()["pages"] == hit.pages
    note = next(r for r in searcher.search("kites wind", k=5) if r.path == "note.md")
    assert (note.page_start, note.page_end, note.pages) == (None, None, None)


def test_pdf_pages_with_small_chunks(library: Path, pdf_extra: None) -> None:
    """Chunks that fit within one page cite exactly that page."""
    sync(library, HashingEmbedder(), max_chars=40)
    hits = Searcher(library, None).search("sourdough starter", k=3)
    hit = next(r for r in hits if r.path == "report.pdf")
    assert (hit.page_start, hit.page_end) == (4, 4)
    assert hit.pages == "p. 4"


def test_docx_headings_become_heading_paths(library: Path, docx_extra: None) -> None:
    sync(library, HashingEmbedder(), max_chars=60)
    hits = Searcher(library, None).search("refuge altitude", k=3)
    hit = next(r for r in hits if r.path == "trip.docx")
    assert hit.heading_path == ["Mountains", "Huts"]
    assert hit.page_start is None


def test_unchanged_documents_are_not_extracted_again(
    library: Path, pdf_extra: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    sync(library, HashingEmbedder())
    pdf = library / "report.pdf"
    pdf.write_bytes(pdf.read_bytes())  # touched: new mtime, same content

    def fail(*_args: Any) -> Any:
        raise AssertionError("extracted again")

    monkeypatch.setattr(documents, "extract", fail)
    stats = sync(library, HashingEmbedder())
    assert stats.up_to_date and not stats.skipped


def test_broken_pdf_is_skipped_by_sync(library: Path, pdf_extra: None) -> None:
    (library / "broken.pdf").write_bytes(b"%PDF-1.4\nnot really")
    (library / "locked.pdf").write_bytes(encrypt_pdf(make_pdf(PDF_PAGES), "secret"))
    stats = sync(library, HashingEmbedder())
    reasons = {s.path: s.reason for s in stats.skipped}
    assert reasons["locked.pdf"] == "encrypted PDF (password required)"
    assert "broken.pdf" in reasons
    assert "report.pdf" not in reasons


def test_cli_search_prints_pages(library: Path, pdf_extra: None) -> None:
    runner = CliRunner()
    result = runner.invoke(app, ["index", str(library), "--embedder", "hashing"])
    assert result.exit_code == 0, result.output
    result = runner.invoke(app, ["search", "glaciers meltwater", "--folder", str(library)])
    assert result.exit_code == 0, result.output
    pdf_line = next(line for line in result.output.splitlines() if "report.pdf:" in line)
    assert "(p. 2)" in pdf_line or "(pp. " in pdf_line


# -- MCP read_context -------------------------------------------------------------
@pytest.fixture
def anyio_backend() -> str:
    return "asyncio"


async def _call(service: NotesService, tool: str, args: dict[str, Any]) -> Any:
    async with Client(create_server(service)) as client:
        result = await client.call_tool(tool, args)
    if result.is_error:
        block = result.content[0]
        assert isinstance(block, TextContent)
        raise AssertionError(block.text)
    return result.structured_content


async def _call_error(service: NotesService, tool: str, args: dict[str, Any]) -> str:
    async with Client(create_server(service)) as client:
        result = await client.call_tool(tool, args)
    assert result.is_error
    block = result.content[0]
    assert isinstance(block, TextContent)
    return block.text


@pytest.mark.anyio
async def test_read_context_serves_pdf_text_with_pages(library: Path, pdf_extra: None) -> None:
    sync(library, HashingEmbedder(), max_chars=40)
    service = NotesService()
    out = await _call(service, "search", {"query": "glaciers retreat", "k": 5})
    hit = next(r for r in out["results"] if r["path"] == "report.pdf")
    assert hit["pages"] == "p. 2" and hit["page_start"] == hit["page_end"] == 2
    ctx = await _call(service, "read_context", {"location": hit["location"], "context": 0})
    assert (ctx["line_start"], ctx["line_end"]) == (hit["line_start"], hit["line_end"])
    assert ctx["pages"] == "p. 2"
    assert "Glaciers retreat" in ctx["text"]
    assert documents.PAGE_BREAK not in ctx["text"]
    whole = await _call(service, "read_context", {"location": "report.pdf"})
    assert (whole["page_start"], whole["page_end"]) == (1, 4)
    assert whole["text"].startswith("Quarterly report")


@pytest.mark.anyio
async def test_read_context_serves_docx_text(library: Path, docx_extra: None) -> None:
    sync(library, HashingEmbedder(), max_chars=60)
    service = NotesService()
    out = await _call(service, "search", {"query": "refuge altitude", "k": 5})
    hit = next(r for r in out["results"] if r["path"] == "trip.docx")
    ctx = await _call(service, "read_context", {"location": hit["location"], "context": 0})
    assert "Refuge | 2800 m" in ctx["text"]
    assert ctx["pages"] is None


@pytest.mark.anyio
async def test_read_context_without_the_extra_refuses_clearly(
    library: Path, pdf_extra: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    sync(library, HashingEmbedder())
    monkeypatch.setitem(sys.modules, "pypdf", None)
    msg = await _call_error(NotesService(), "read_context", {"location": "report.pdf:1-2"})
    assert msg.endswith(
        "cannot read report.pdf: PDF support is not installed; install notelight[pdf]"
    )


@pytest.mark.anyio
async def test_read_context_extraction_is_cached(
    library: Path, pdf_extra: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    sync(library, HashingEmbedder())
    service = NotesService()
    calls = []
    real = documents.load_text

    def counting(suffix: str, data: bytes) -> Any:
        calls.append(suffix)
        return real(suffix, data)

    monkeypatch.setattr(documents, "load_text", counting)
    for _ in range(3):
        await _call(service, "read_context", {"location": "report.pdf:1"})
    assert calls == [".pdf"]
    (library / "report.pdf").write_bytes(make_pdf([["Rewritten page."]]))
    out = await _call(service, "read_context", {"location": "report.pdf:1"})
    assert out["text"] == "Rewritten page."
    assert calls == [".pdf", ".pdf"]
